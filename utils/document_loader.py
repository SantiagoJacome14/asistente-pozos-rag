"""
Módulo para cargar y procesar documentos en diferentes formatos.
Soporta: PDF, DOCX, XLSX
"""

import os
from typing import List, Tuple
import PyPDF2
from docx import Document
import openpyxl
import pandas as pd
import logging

logger = logging.getLogger(__name__)


class DocumentLoader:
    """Cargador de documentos multiformato."""

    def __init__(self):
        self.supported_formats = ['.pdf', '.docx', '.xlsx', '.xls']

    def load_document(self, file_path: str) -> Tuple[str, str]:
        """
        Carga un documento y extrae su contenido de texto.
        
        Args:
            file_path: Ruta al archivo
            
        Returns:
            Tupla (contenido_texto, nombre_archivo)
            
        Raises:
            ValueError: Si el formato no es soportado
            Exception: Si hay error al procesar el archivo
        """
        file_ext = os.path.splitext(file_path)[1].lower()
        file_name = os.path.basename(file_path)

        if file_ext not in self.supported_formats:
            raise ValueError(
                f"Formato no soportado: {file_ext}. "
                f"Formatos válidos: {self.supported_formats}"
            )

        try:
            if file_ext == '.pdf':
                content = self._load_pdf(file_path)
            elif file_ext == '.docx':
                content = self._load_docx(file_path)
            elif file_ext in ['.xlsx', '.xls']:
                content = self._load_excel(file_path)

            logger.info(f"Documento cargado exitosamente: {file_name}")
            return content, file_name

        except Exception as e:
            logger.error(f"Error cargando documento {file_name}: {str(e)}")
            raise

    def _load_pdf(self, file_path: str) -> str:
        """Extrae texto de un PDF."""
        content = []
        
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                num_pages = len(pdf_reader.pages)
                
                for page_num in range(num_pages):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()
                    if text.strip():
                        content.append(f"[Página {page_num + 1}]\n{text}")
                        
        except Exception as e:
            raise Exception(f"Error al procesar PDF: {str(e)}")

        return "\n\n".join(content) if content else ""

    def _load_docx(self, file_path: str) -> str:
        """Extrae texto de un DOCX."""
        content = []
        
        try:
            doc = Document(file_path)
            
            for para in doc.paragraphs:
                if para.text.strip():
                    content.append(para.text)
            
            # Procesar tablas
            for table in doc.tables:
                content.append("\n[TABLA]\n")
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    content.append(" | ".join(row_data))
                    
        except Exception as e:
            raise Exception(f"Error al procesar DOCX: {str(e)}")

        return "\n".join(content) if content else ""

    def _load_excel(self, file_path: str) -> str:
        """Extrae contenido de un archivo Excel."""
        content = []
        
        try:
            excel_file = pd.ExcelFile(file_path)
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                content.append(f"[HOJA: {sheet_name}]\n")
                content.append(df.to_string())
                content.append("\n")
                    
        except Exception as e:
            raise Exception(f"Error al procesar Excel: {str(e)}")

        return "\n".join(content) if content else ""

    def validate_file(self, file_path: str) -> bool:
        """Valida que el archivo exista y sea de formato soportado."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Archivo no encontrado: {file_path}")
        
        file_ext = os.path.splitext(file_path)[1].lower()
        if file_ext not in self.supported_formats:
            return False
        
        return True
